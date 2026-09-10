"""
Loads the Variety / VarietyGroup / IANNumbers / PUC / Commodity reference
tables (used for cross-checking scanned label values, e.g. PUC -> GGN) into
DynamoDB from WhatsAppTables.xlsx. Idempotent/safe to re-run whenever Dole
issues updated reference data -- same pattern as sync_spec_catalog.py.

Usage:
    python tools/sync_reference_tables.py [path-to-xlsx]
"""
import re
import sys
import os

import boto3
import openpyxl
from botocore.exceptions import ClientError


def normalize_key_text(value):
    """
    Collapses whitespace runs to a single space before upper-casing, matching
    the normalization the webhook applies to values it transcribes off a
    label (which can carry a line-wrap newline where this source data has a
    plain space) -- keeps stored keys and lookup keys on the same footing.
    """
    return re.sub(r"\s+", " ", str(value)).strip().upper()

REGION = "us-east-2"
DEFAULT_XLSX_PATH = os.path.join(os.path.dirname(__file__), "..", "WhatsAppTables.xlsx")
INDIA_ADDRESSES_XLSX_PATH = os.path.join(os.path.dirname(__file__), "extracted_indian_addresses.xlsx")
AGENTS_XLSX_PATH = os.path.join(os.path.dirname(__file__), "Agents.xlsx")

VARIETY_TABLE = "whatsapp-variety"
VARIETY_GROUP_TABLE = "whatsapp-variety-group"
IAN_NUMBERS_TABLE = "whatsapp-ian-numbers"
PUC_TABLE = "whatsapp-puc"
COMMODITY_TABLE = "whatsapp-commodity"
REWE_COMBINATIONS_TABLE = "whatsapp-rewe-combinations"
AGENT_ADDRESSES_TABLE = "whatsapp-agent-addresses"

# Agent/importer codes are unique across every country that uses this table
# (confirmed by the user), so one shared table + one flat "Code" key covers
# all of them -- no need for a per-country namespace. Non-India rows come
# from tools/Agents.xlsx (see sync_general_agents()) rather than being
# hardcoded, now that a real source spreadsheet exists for them.
#
# One genuine cross-country code collision was found and resolved by the
# user 2026-09-10: code "XX" existed in the India sheet as a duplicate of
# "TA" (same company "SAMRIDDHI TRADE", same FSSAI -- a data-entry error in
# the original India source, not a real second consignee) and *separately*
# in Agents.xlsx as a real, different company ("PT. PARAHYANGAN SUKSES
# FOODS", Indonesia). Resolution: the India loader below explicitly skips
# "XX" (keeping only "TA" for Samriddhi Trade); Agents.xlsx's "XX" is the
# real occupant of that code going forward.
INDIA_DUPLICATE_CODES_TO_SKIP = {"XX"}

# REWE (13A/13B/13C/13D) rule 6: which GTIN/supplier-code/bilingual-name/
# packaging text is correct depends on Pack + VarietyGroup, except the
# BX/RX/3X variety groups which also depend on InventoryCode. Rather than
# teach the lookup engine that bucketing logic, every simple combination is
# stored under a 2-part key ("Pack#VarietyGroup") and only the one
# InvCode='RH' exception gets a 3-part key ("Pack#VarietyGroup#InvCode") --
# the webhook tries the 3-part key first, falling back to the 2-part one
# (see lambda_function.py's lookup_rules "fallback_key"). Hardcoded directly
# from the user's rules (no spreadsheet source), same pattern as
# sync_spec_catalog.py's embedded SPEC_ROWS.
REWE_COMBINATION_ROWS = []


def _add_rewe_rows(packs, variety_group, nan_puc, gtin, de, en, packaging, inv_code=None):
    for pack in packs:
        key_parts = [pack, variety_group] + ([inv_code] if inv_code else [])
        REWE_COMBINATION_ROWS.append({
            "lookup_key": "#".join(key_parts),
            "ReweNanPUC": nan_puc,
            "GTIN": gtin,
            "VarietyGroupDE": de,
            "VarietyGroupEN": en,
            "REWEPackaging": packaging,
        })


# Pack A04I / B04I bucket (4.5kg)
_add_rewe_rows(["A04I", "B04I"], "WS", "7722245 - D7520", "4337256546317", "Trauben Hell Kernlos", "White Seedless", "4.5kg")
_add_rewe_rows(["A04I", "B04I"], "RS", "8873554 - D7520", "4337256562003", "Trauben Rot Kernlos", "Red Seedless", "4.5kg")
_add_rewe_rows(["A04I", "B04I"], "BS", "8952160 - D7520", "4337256558372", "Trauben Blau Kernlos", "Black Seedless", "4.5kg")

# Pack A05D / A05A bucket (10x500g)
_add_rewe_rows(["A05D", "A05A"], "WS", "7412355 - D7520", "4046118840691", "Trauben Hell Kernlos", "White Seedless", "10x500g")
_add_rewe_rows(["A05D", "A05A"], "RS", "9232885 - D7520", "4051826535715", "Trauben Dunkel Kernlos", "Red Seedless", "10x500g")
_add_rewe_rows(["A05D", "A05A"], "BS", "9232885 - D7520", "4051826535715", "Trauben Dunkel Kernlos", "Black Seedless", "10x500g")
# BX/RX/3X default (InventoryCode != 'RH')
for _vg in ("BX", "RX", "3X"):
    _add_rewe_rows(["A05D", "A05A"], _vg, "377024 - D7520", "4051826598116", "Trauben Mix Kernlos", "Mixed Varieties", "10x500g")
# BX/RX/3X exception (InventoryCode == 'RH')
for _vg in ("BX", "RX", "3X"):
    _add_rewe_rows(
        ["A05D", "A05A"], _vg, "9001853 - D7520", "4337256718363",
        "Rewe Beste Wahl Tafeltrauben Mix, kernlos", "Mixed Varieties", "10x500g", inv_code="RH",
    )

dynamodb = boto3.resource("dynamodb", region_name=REGION)


def ensure_table(table_name, key_schema, attribute_definitions):
    client = dynamodb.meta.client
    try:
        client.describe_table(TableName=table_name)
        print(f"Table {table_name} already exists.")
        return
    except ClientError as e:
        if e.response["Error"]["Code"] != "ResourceNotFoundException":
            raise

    print(f"Creating table {table_name}...")
    client.create_table(
        TableName=table_name,
        KeySchema=key_schema,
        AttributeDefinitions=attribute_definitions,
        BillingMode="PAY_PER_REQUEST",
    )
    client.get_waiter("table_exists").wait(TableName=table_name)
    print(f"Table {table_name} created.")


def rows(ws):
    """Yields each data row (after the header) as a dict keyed by header cell."""
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        yield dict(zip(header, row))


def sync_variety(wb):
    # VarietyName isn't globally unique on its own (a handful of names are
    # reused across different commodities, e.g. "EUREKA" is both a Blueberry
    # and a Lemon variety) -- Commodity is the range key so both rows survive
    # instead of one silently overwriting the other.
    ensure_table(
        VARIETY_TABLE,
        key_schema=[
            {"AttributeName": "VarietyName", "KeyType": "HASH"},
            {"AttributeName": "Commodity", "KeyType": "RANGE"},
        ],
        attribute_definitions=[
            {"AttributeName": "VarietyName", "AttributeType": "S"},
            {"AttributeName": "Commodity", "AttributeType": "S"},
        ],
    )
    table = dynamodb.Table(VARIETY_TABLE)
    count = 0
    with table.batch_writer(overwrite_by_pkeys=["VarietyName", "Commodity"]) as batch:
        for r in rows(wb["Variety"]):
            batch.put_item(Item={
                "VarietyName": normalize_key_text(r["VarietyName"]),
                "Commodity": normalize_key_text(r["Commodity"]),
                "VarietyCode": str(r["VarietyCode"]).strip(),
                "VarietyGroupCode": str(r["VarietyGroupCode"]).strip(),
            })
            count += 1
    print(f"{VARIETY_TABLE}: loaded {count} rows.")


def sync_variety_group(wb):
    # VarietyGroupCode is confirmed globally unique (no two commodities reuse
    # the same code), so it's a clean single hash key.
    ensure_table(
        VARIETY_GROUP_TABLE,
        key_schema=[{"AttributeName": "VarietyGroupCode", "KeyType": "HASH"}],
        attribute_definitions=[{"AttributeName": "VarietyGroupCode", "AttributeType": "S"}],
    )
    table = dynamodb.Table(VARIETY_GROUP_TABLE)
    count = 0
    with table.batch_writer(overwrite_by_pkeys=["VarietyGroupCode"]) as batch:
        for r in rows(wb["VarietyGroup"]):
            item = {
                "VarietyGroupCode": str(r["VarietyGroupCode"]).strip(),
                "Commodity": str(r["Commodity"]).strip().upper(),
            }
            if r.get("LidlTranslations"):
                item["LidlTranslations"] = str(r["LidlTranslations"]).strip()
            batch.put_item(Item=item)
            count += 1
    print(f"{VARIETY_GROUP_TABLE}: loaded {count} rows.")


def sync_ian_numbers(wb):
    # Four columns (Commodity, VarietyGroup, Pack, InventoryCode) make up the
    # real key -- DynamoDB only supports a hash+range key pair, so they're
    # concatenated into a single hash key rather than modeling a composite
    # range key for what's ultimately just an exact-match lookup.
    ensure_table(
        IAN_NUMBERS_TABLE,
        key_schema=[{"AttributeName": "lookup_key", "KeyType": "HASH"}],
        attribute_definitions=[{"AttributeName": "lookup_key", "AttributeType": "S"}],
    )
    table = dynamodb.Table(IAN_NUMBERS_TABLE)
    count = 0
    with table.batch_writer(overwrite_by_pkeys=["lookup_key"]) as batch:
        for r in rows(wb["IANNumbers"]):
            commodity = str(r["Commodity"]).strip().upper()
            variety_group = str(r["VarietyGroup"]).strip().upper()
            pack = str(r["Pack"]).strip().upper()
            inventory_code = str(r["InventoryCode"]).strip().upper()
            batch.put_item(Item={
                "lookup_key": f"{commodity}#{variety_group}#{pack}#{inventory_code}",
                "Commodity": commodity,
                "VarietyGroup": variety_group,
                "Pack": pack,
                "InventoryCode": inventory_code,
                "IANNumber": str(int(r["IANNumber"])),
            })
            count += 1
    print(f"{IAN_NUMBERS_TABLE}: loaded {count} rows.")


def sync_puc(wb):
    ensure_table(
        PUC_TABLE,
        key_schema=[{"AttributeName": "PUC", "KeyType": "HASH"}],
        attribute_definitions=[{"AttributeName": "PUC", "AttributeType": "S"}],
    )
    table = dynamodb.Table(PUC_TABLE)
    count = 0
    with table.batch_writer(overwrite_by_pkeys=["PUC"]) as batch:
        for r in rows(wb["PUC"]):
            batch.put_item(Item={
                "PUC": str(r["PUC"]).strip().upper(),
                "GGN": str(r["GGN"]).strip(),
            })
            count += 1
    print(f"{PUC_TABLE}: loaded {count} rows.")


def sync_commodity(wb):
    ensure_table(
        COMMODITY_TABLE,
        key_schema=[{"AttributeName": "Commodity", "KeyType": "HASH"}],
        attribute_definitions=[{"AttributeName": "Commodity", "AttributeType": "S"}],
    )
    table = dynamodb.Table(COMMODITY_TABLE)
    count = 0
    with table.batch_writer(overwrite_by_pkeys=["Commodity"]) as batch:
        for r in rows(wb["Commodity"]):
            batch.put_item(Item={
                "Commodity": str(r["Commodity"]).strip().upper(),
                "CommodityName": str(r["CommodityName"]).strip(),
                # The source spreadsheet's header calls this "IANNumber" but
                # confirmed with the user this is actually the Lidl Product
                # Number (17-char code), not a real IAN number.
                "LidlProductNumber": str(r["IANNumber"]).strip(),
            })
            count += 1
    print(f"{COMMODITY_TABLE}: loaded {count} rows.")


def sync_rewe_combinations():
    ensure_table(
        REWE_COMBINATIONS_TABLE,
        key_schema=[{"AttributeName": "lookup_key", "KeyType": "HASH"}],
        attribute_definitions=[{"AttributeName": "lookup_key", "AttributeType": "S"}],
    )
    table = dynamodb.Table(REWE_COMBINATIONS_TABLE)
    count = 0
    with table.batch_writer(overwrite_by_pkeys=["lookup_key"]) as batch:
        for row in REWE_COMBINATION_ROWS:
            batch.put_item(Item=row)
            count += 1
    print(f"{REWE_COMBINATIONS_TABLE}: loaded {count} rows.")


def split_name_and_address(address1, *rest_lines):
    """
    Some source sheets run the company's full legal name straight into the
    first address line with no delimiter but a comma (e.g. "TENGINE SDN BHD,
    LOT 2457," or "SAHAJ LAXMI TRADERS, SHOP NO 23, GROUND FLOOR,"). Splitting
    at the first comma recovers a clean name for those rows -- the address
    fragment left over (if any) joins the rest of the address lines instead
    of staying glued to the name. Rows with no comma at all (Address1 is
    already just the name) are returned whole, unchanged. This matters
    because matching a name against the *whole* Address1 (name+address)
    would score low for a genuinely correct short name against a long
    combined string -- exactly the false-failure this split avoids. Shared
    by both the India and general-agents loaders below (same source-data
    shape, same fix).
    """
    name_part, _, addr1_remainder = address1.partition(",")
    importer_name = name_part.strip()
    address_lines = " ".join(p for p in [addr1_remainder.strip(), *rest_lines] if p)
    return importer_name, address_lines


def sync_agent_addresses(xlsx_path=INDIA_ADDRESSES_XLSX_PATH):
    # Started as India-only (17C's consignee/agent code -> name, 3-line
    # address, FSSAI license number), generalized 2026-09-10 once 17A
    # (Malaysia) needed the exact same shape (agent code -> importer name +
    # address) -- renamed from whatsapp-india-addresses to this generic name.
    # The India rows come from a separate xlsx (not WhatsAppTables.xlsx,
    # since it's not Dole-global); sync_general_agents() below loads the
    # non-India rows from a second, separate sheet into the same table.
    #
    # Only ImporterName/AddressLines (the split results) and FSSAI are
    # actually referenced by any validation config -- the raw Address1/
    # Address2/Address3/FullAddress components aren't kept in the table at
    # all, deliberately: this table is hand-editable via the dashboard's
    # Reference Tables tab, and keeping unused raw fields alongside the
    # derived ones they were split from would let someone edit one without
    # the other, silently drifting the two out of sync.
    ensure_table(
        AGENT_ADDRESSES_TABLE,
        key_schema=[{"AttributeName": "Code", "KeyType": "HASH"}],
        attribute_definitions=[{"AttributeName": "Code", "AttributeType": "S"}],
    )
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    table = dynamodb.Table(AGENT_ADDRESSES_TABLE)
    count = 0
    with table.batch_writer(overwrite_by_pkeys=["Code"]) as batch:
        for r in rows(wb["Addresses"]):
            code = normalize_key_text(r["Code"])
            if code in INDIA_DUPLICATE_CODES_TO_SKIP:
                continue
            address2 = str(r["Address2"]).strip()
            address3 = str(r["Inv: IndianAddress3"]).strip()
            importer_name, address_lines = split_name_and_address(str(r["Address1"]).strip(), address2, address3)
            item = {
                "Code": code,
                "ImporterName": importer_name,
                "AddressLines": address_lines,
            }
            if r.get("FSSAI"):
                item["FSSAI"] = str(r["FSSAI"]).strip()
            batch.put_item(Item=item)
            count += 1
    print(f"{AGENT_ADDRESSES_TABLE}: loaded {count} India rows.")


def sync_general_agents(xlsx_path=AGENTS_XLSX_PATH):
    # Non-India agent/importer codes (Malaysia's 17A and onward) -- same
    # table, same Code/ImporterName/AddressLines shape as the India loader
    # above, just a different source sheet (Code/Customer/Address1-4, no
    # FSSAI). "Customer" is a short internal nickname (e.g. "CHOP TONG"),
    # not what's printed on a real label -- Address1 carries the full legal
    # company name instead (confirmed against a real 17A photo: agent "5K"'s
    # Address1 "T&T FRESH SDN BHD" is exactly what the label prints), so
    # Address1 feeds ImporterName via the same split_name_and_address() used
    # for India, not the Customer column.
    ensure_table(
        AGENT_ADDRESSES_TABLE,
        key_schema=[{"AttributeName": "Code", "KeyType": "HASH"}],
        attribute_definitions=[{"AttributeName": "Code", "AttributeType": "S"}],
    )
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    table = dynamodb.Table(AGENT_ADDRESSES_TABLE)
    count = 0
    with table.batch_writer(overwrite_by_pkeys=["Code"]) as batch:
        for r in rows(wb["Agents"]):
            address2 = str(r.get("Address2") or "").strip()
            address3 = str(r.get("Address3") or "").strip()
            address4 = str(r.get("Address4") or "").strip()
            importer_name, address_lines = split_name_and_address(
                str(r["Address1"]).strip(), address2, address3, address4
            )
            batch.put_item(Item={
                "Code": normalize_key_text(r["Code"]),
                "ImporterName": importer_name,
                "AddressLines": address_lines,
            })
            count += 1
    print(f"{AGENT_ADDRESSES_TABLE}: loaded {count} general-agent rows.")


def main():
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX_PATH
    print(f"Loading {xlsx_path} ...")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    sync_variety(wb)
    sync_variety_group(wb)
    sync_ian_numbers(wb)
    sync_puc(wb)
    sync_commodity(wb)
    sync_rewe_combinations()
    sync_agent_addresses()
    sync_general_agents()


if __name__ == "__main__":
    main()
